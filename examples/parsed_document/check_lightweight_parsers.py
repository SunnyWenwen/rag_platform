"""Run real local MarkItDown / PyMuPDF4LLM conversions on owned fixtures.

No HTTP service, Engine, OCR, cloud model or Chunker is executed. Generated
binary files / full normalized outputs stay in --work-dir, outside the repo.
"""
from collections import Counter
from copy import deepcopy
from io import BytesIO
import argparse
import base64
import hashlib
import importlib.metadata
import json
from pathlib import Path
import sys

from jsonschema import Draft202012Validator
import pymupdf
import pymupdf4llm
from markitdown import MarkItDown, StreamInfo

from markdown_parser_adapters import from_markdown, from_page_chunks, from_layout_json
from validate import validate_document, TEXT_SPANS, PAGE_BLOCKS

ROOT = Path(__file__).resolve().parents[2]


def portable(value):
    return json.loads(json.dumps(value, ensure_ascii=False, allow_nan=False))


def make_fixtures(work):
    work.mkdir(parents=True, exist_ok=True)
    pdf = work / 'maintenance.pdf'
    doc = pymupdf.open()
    for index in range(3):
        page = doc.new_page(width=595, height=842)
        if index == 0:
            page.insert_text((72, 70), 'Maintenance Guide', fontsize=20, fontname='hebo')
            page.insert_text((72, 108), 'Disconnect power before cleaning.', fontsize=11)
            xs, ys = [72, 260, 440], [150, 176, 202]
            for x in xs:
                page.draw_line((x, ys[0]), (x, ys[-1]))
            for y in ys:
                page.draw_line((xs[0], y), (xs[-1], y))
            for x, y, text in [(80, 168, 'Item'), (268, 168, 'Frequency'), (80, 194, 'Fan'), (268, 194, 'Weekly')]:
                page.insert_text((x, y), text, fontsize=11)
        elif index == 1:
            page.insert_text((72, 70), 'Fan cleaning', fontsize=18, fontname='hebo')
            page.insert_text((72, 108), 'Clean the fan every week.', fontsize=11)
            page.insert_text((72, 145), "print('inspection complete')", fontsize=10, fontname='cour')
        # Page 3 is intentionally blank.
    doc.set_metadata({'title': 'Maintenance Guide', 'author': 'RAG Builder contract fixture',
                      'creationDate': 'D:20261030000000Z', 'modDate': 'D:20261030000000Z'})
    doc.save(pdf, no_new_id=True)
    doc.close()
    from docx import Document
    word = Document()
    word.add_heading('Maintenance Guide', level=1)
    word.add_paragraph('Disconnect power before cleaning.')
    table = word.add_table(rows=2, cols=2)
    for cell, text in zip([c for r in table.rows for c in r.cells], ['Item', 'Frequency', 'Fan', 'Weekly']):
        cell.text = text
    word.save(work / 'maintenance.docx')
    (work / 'maintenance.html').write_text('<html><head><title>Maintenance Guide</title></head><body><h1>Maintenance Guide</h1><p>Disconnect power before cleaning.</p><table><tr><th>Item</th><th>Frequency</th></tr><tr><td>Fan</td><td>Weekly</td></tr></table></body></html>')
    return {ext: work / ('maintenance.' + ext) for ext in ('pdf', 'docx', 'html')}


def run(work):
    files = make_fixtures(work)
    ingestion_validator = Draft202012Validator(json.loads((ROOT / 'schemas/ingestion_document.schema.json').read_text()))
    results, calls, inputs = [], [], []

    def payload_bytes(path, media_type):
        original = path.read_bytes()
        revision = hashlib.sha256(original).hexdigest()
        payload = {'document_id': path.stem, 'source_revision': revision,
                   'base64': base64.b64encode(original).decode(), 'media_type': media_type}
        ingestion_validator.validate(payload)
        decoded = base64.b64decode(payload['base64'], validate=True)
        assert decoded == original
        inputs.append({'file': path.name, 'media_type': media_type, 'sha256': revision,
                       'base64_ingestion_schema_valid': True, 'decoded_bytes_equal': True})
        return decoded, revision

    def check(name, document, native):
        validate_document(document)
        assert document['metadata']['_rag']['native_result'] == portable(native)
        structure = document['structure'].get('text')
        blocks = [] if not structure else (structure['blocks'] if 'blocks' in structure else [b for p in structure.get('pages', []) for b in p['blocks']])
        item = {'id': name, 'schema_and_semantics_valid': True, 'native_json_result_retained': True,
                'profile': structure['schema_id'] if structure else 'unstructured',
                'block_types': dict(Counter(b['type'] for b in blocks)),
                'text_chars': len(document['fields']['text']),
                'has_page_locations': any(l['kind'] == 'page' for b in blocks for l in b.get('locations', []))}
        if structure and 'pages' in structure:
            item['page_numbers'] = [p['page_number'] for p in structure['pages']]
        results.append(item)
        (work / (name + '.parsed.json')).write_text(json.dumps(document, ensure_ascii=False, indent=2) + '\n')

    converter = MarkItDown(enable_plugins=False)
    media_types = {'pdf': 'application/pdf', 'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'html': 'text/html'}
    for ext, path in files.items():
        data, revision = payload_bytes(path, media_types[ext])
        value = converter.convert_stream(BytesIO(data), stream_info=StreamInfo(extension='.'+ext, mimetype=media_types[ext], filename=path.name))
        native = {'markdown': value.markdown, 'title': value.title}
        assert value.text_content == value.markdown
        calls.append({'provider': 'markitdown', 'input': ext, 'output': 'DocumentConverterResult',
                      'keys': list(vars(value)), 'pages_attribute': hasattr(value, 'pages'),
                      'markdown_table_present': '| ---' in value.markdown or '|---' in value.markdown})
        for structured in (False, True):
            name = f'markitdown-{ext}-' + ('spans' if structured else 'plain')
            document = from_markdown(value.markdown, document_id=name, revision=revision, native=native,
                                     provider='markitdown', structured=structured, title=value.title)
            assert document['fields']['text'] == value.markdown
            check(name, document, native)

    data, revision = payload_bytes(files['pdf'], 'application/pdf')
    with pymupdf.open(stream=data, filetype='pdf') as source:
        pymupdf.use_layout(False)
        native = source.to_markdown(show_progress=False)
        calls.append({'provider': 'pymupdf4llm', 'mode': 'legacy-markdown', 'output': 'str', 'ocr': False})
        for structured in (False, True):
            name = 'pymupdf-legacy-markdown-' + ('spans' if structured else 'plain')
            document = from_markdown(native, document_id=name, revision=revision, native=native,
                                     provider='pymupdf4llm-legacy', structured=structured)
            assert document['fields']['text'] == native
            check(name, document, native)
        try:
            source.to_json(show_progress=False)
        except NotImplementedError:
            legacy_json_requires_layout = True
        else:
            raise AssertionError('pinned mode assumption changed: review JSON compatibility')
        for layout in (False, True):
            pymupdf.use_layout(layout)
            kwargs = {'use_ocr': False} if layout else {}
            native = portable(source.to_markdown(page_chunks=True, show_progress=False, **kwargs))
            mode = 'layout-pages' if layout else 'legacy-pages'
            keys = list(native[0])
            calls.append({'provider': 'pymupdf4llm', 'mode': mode, 'output': 'page list', 'ocr': False,
                          'page_dictionary_keys': keys, 'page_metadata_keys': list(native[0]['metadata'])})
            for profile in (TEXT_SPANS, PAGE_BLOCKS):
                name = 'pymupdf-' + mode + '-' + profile.split('/')[0]
                document = from_page_chunks(native, document_id=name, revision=revision,
                                            provider='pymupdf4llm-'+mode, profile=profile)
                check(name, document, native)
                assert results[-1]['page_numbers'] == [1, 2, 3]
                if profile == PAGE_BLOCKS:
                    assert document['structure']['text']['pages'][-1]['blocks'] == []
        output = source.to_json(use_ocr=False, show_progress=False)
        native = json.loads(output) if isinstance(output, str) else portable(output)
        calls.append({'provider': 'pymupdf4llm', 'mode': 'layout-json', 'output': 'JSON', 'ocr': False,
                      'box_classes': dict(Counter(b['boxclass'] for p in native['pages'] for b in p['boxes'])),
                      'table_column_counts': [b['table']['col_count'] for p in native['pages'] for b in p['boxes'] if b['boxclass'] == 'table']})
        for profile in (TEXT_SPANS, PAGE_BLOCKS):
            name = 'pymupdf-layout-json-' + profile.split('/')[0]
            document = from_layout_json(native, document_id=name, revision=revision, profile=profile)
            check(name, document, native)
            structure = document['structure']['text']
            mapped = structure['blocks'] if 'blocks' in structure else [b for p in structure.get('pages', []) for b in p['blocks']]
            assert len(mapped) == sum(len(p['boxes']) for p in native['pages'])
            for block, box in zip(mapped, [b for p in native['pages'] for b in p['boxes']]):
                assert block['extensions']['source.pymupdf'] == box
                if block['type'] == 'table':
                    assert document['fields'][block['data_ref']]['payload'] == box['table']
        alias = pymupdf4llm.to_markdown(source, pages=[0], use_ocr=False, show_progress=False)
        preferred = source.to_markdown(pages=[0], use_ocr=False, show_progress=False)
        assert alias == preferred

    # Explicit failure cases: page identity and recognized structured output.
    negative = 0
    for bad in [[], [{'metadata': {}, 'text': 'text'}], [{'metadata': {'page': 1, 'page_number': 2}, 'text': 'text'}]]:
        try:
            from_page_chunks(bad, document_id='bad', revision='v1', provider='test')
        except ValueError:
            negative += 1
        else:
            raise AssertionError('invalid page data accepted')
    bad_layout = deepcopy(native)
    bad_layout['pages'][0]['boxes'][0]['boxclass'] = 'unknown-test-class'
    try:
        from_layout_json(bad_layout, document_id='bad', revision='v1')
    except ValueError:
        negative += 1
    else:
        raise AssertionError('unknown native class silently accepted')
    versions = {p: importlib.metadata.version(p) for p in ('markitdown', 'pymupdf', 'pymupdf4llm', 'pymupdf-layout', 'markdown-it-py', 'jsonschema', 'python-docx')}
    return {'date': '2026-10-30', 'contract': 'parsed-document/0.1-draft', 'versions': versions,
            'scope': 'Actual local synthetic PDF/DOCX/HTML conversions from validated Base64 ingestion bytes; no HTTP Engine, external API, OCR, image assets, real-world corpus or downstream Chunker.',
            'fixture': {'owned_synthetic_inputs': True, 'pdf_pages': 3, 'blank_page_number': 3, 'expected_table_columns': 2},
            'inputs': inputs, 'successful_parser_calls': len(calls), 'parser_calls': calls,
            'validated_mappings': len(results), 'negative_adapter_cases': negative,
            'legacy_json_requires_layout': legacy_json_requires_layout,
            'document_method_equals_compatibility_alias': True, 'results': results}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--work-dir', type=Path, required=True)
    parser.add_argument('--report', type=Path, required=True)
    args = parser.parse_args()
    result = run(args.work_dir)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'successful_parser_calls': result['successful_parser_calls'], 'validated_mappings': result['validated_mappings'],
                      'negative_adapter_cases': result['negative_adapter_cases'], 'report': str(args.report)}, ensure_ascii=False))
